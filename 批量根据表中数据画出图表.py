import openpyxl
from openpyxl.chart import Reference,BarChart
import os
# 批量根据表中数据画出图表
for file_name in os.listdir('./test-data2'):
    # excel文件的完整路径 ./test-data2/1.xlsx
    file_name = os.path.join('./test-data2',file_name)
    ex_file = openpyxl.load_workbook(file_name)
    sheet_names =  ex_file.sheetnames
    for sheet_name in sheet_names:
        sheet_file = ex_file[sheet_name]
        # 01 创建一个reference对象 表示作用在图表中的数据区域
        data = Reference(sheet_file,min_col=2,min_row=1,max_col=3,max_row=13)
        # 02 创建图表对象
        lc = BarChart()
        lc.title = sheet_name
        lc.x_axis.title = '日期'
        lc.y_axis.title = '营收额'
        # 03 向图表对象中添加数据
        lc.add_data(data,titles_from_data=True)
        # 04 使用日期作为这一列的x轴
        x_label = Reference(sheet_file,min_col=1,min_row=2,max_row=13)
        lc.set_categories(x_label)
        # 05 将图表添加到制定的sheet中
        sheet_file.add_chart(lc,'E5')
        ex_file.save(file_name)

